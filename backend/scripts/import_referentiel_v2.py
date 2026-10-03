#!/usr/bin/env python3
# 2026-10-02 — import v2 : fichiers Excel réels (un .xlsx par table, 1re feuille,
# SANS ligne d'en-tête, colonnes en ordre fixe — voir SPECS ci-dessous).
# Les clés étrangères sont des CODES (DUPL, 9AL, AM...) ou des NUMEROS legacy
# (criteres 1..40, details 1..85, profils 1..8) : le script fait la correspondance.
# Usage : python scripts/import_referentiel_v2.py /chemin/vers/dossier_excel
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, ".")

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.core.security import hash_password
from app.models.rating import (
    Categorie, Critere, CritereDetail, Departement, Emploi, EmploiProfil,
    Poste, Profil, ProfilCritere, Salarie, Section, Site,
)

# Familles d'emplois déduites du code (1 Cadres / 2 AM / 3 Empl.-Ouvr.) — à ajuster si besoin
FAMILLE_PAR_CODE = {"CD": 1, "CM": 1, "CS": 1, "AM": 2,
                    "EN": 3, "EQ": 3, "ON": 3, "OQ": 3}

# Ordre des colonnes attendu dans chaque feuille (fichier = nom.xlsx)
SPECS = {
    "Sites":           ["code", "libelle"],
    "Departements":    ["code", "libelle", "site_code"],
    "Sections":        ["code", "libelle", "dept_code"],
    "Emplois":         ["code", "libelle"],
    "Categories":      ["code"],
    "Postes":          ["code", "libelle"],
    "Criteres":        ["num", "libelle"],
    "CritereDetails":  ["num", "critere_num", "libelle", "valeur"],
    "Profils":         ["num", "libelle"],
    "ProfilCriteres":  ["profil_num", "critere_num", "coefficient", "ordre"],
    "EmploiProfils":   ["emploi_code", "profil_num", "ordre"],
    "Salaries":        ["matricule", "nom", "prenoms", "site_code", "dept_code",
                        "section_code", "emploi_code", "categorie_code", "poste_code",
                        "date_embauche", "email", "n1_matricule", "hors_evaluation"],
}


def lire(fichier: Path):
    """1re feuille, sans en-tête -> liste de dicts selon SPECS (valeurs nettoyées)."""
    spec = SPECS[fichier.stem]
    ws = load_workbook(fichier, read_only=True, data_only=True).worksheets[0]
    lignes = []
    for row in ws.iter_rows(values_only=True):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        d = {}
        for i, col in enumerate(spec):
            v = row[i] if i < len(row) else None
            if isinstance(v, str):
                v = v.strip()
                if v == "":
                    v = None
            d[col] = v
        lignes.append(d)
    return lignes


def en_bool(v) -> bool:
    """Convertit 1/0, Oui/Non, True/False en booléen (robuste aux exports Excel)."""
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    return s in ("1", "true", "vrai", "oui", "yes", "y")


def main() -> None:
    if len(sys.argv) != 2:
        print("Usage : import_referentiel_v2.py /chemin/vers/dossier_excel")
        sys.exit(1)
    dossier = Path(sys.argv[1])
    manquants = [f + ".xlsx" for f in SPECS if not (dossier / (f + ".xlsx")).exists()]
    if manquants:
        print("Fichiers manquants dans " + str(dossier) + " : " + ", ".join(manquants))
        sys.exit(1)

    db: Session = SessionLocal()

    # --- Sites / Départements / Sections (FK par code) ---
    sites: dict[str, int] = {}
    for r in lire(dossier / "Sites.xlsx"):
        s = Site(code=str(r["code"]), libelle=str(r["libelle"]))
        db.add(s); db.flush()
        sites[s.code] = s.id
    print("Sites :", len(sites))

    deps: dict[str, int] = {}
    for r in lire(dossier / "Departements.xlsx"):
        site_id = sites.get(r["site_code"]) if r["site_code"] else None
        if r["site_code"] and not site_id:
            print("  ! département " + str(r["code"]) + " : site '" + str(r["site_code"]) + "' inconnu -> NULL")
        d = Departement(code=str(r["code"]), libelle=str(r["libelle"]), site_id=site_id)
        db.add(d); db.flush()
        deps[d.code] = d.id
    print("Departements :", len(deps))

    secs: dict[str, int] = {}
    for r in lire(dossier / "Sections.xlsx"):
        dep_id = deps.get(r["dept_code"]) if r["dept_code"] else None
        if r["dept_code"] and not dep_id:
            print("  ! section " + str(r["code"]) + " : département '" + str(r["dept_code"]) + "' inconnu")
            continue
        s = Section(code=str(r["code"]), libelle=str(r["libelle"]), departement_id=dep_id)
        db.add(s); db.flush()
        secs[s.code] = s.id
    print("Sections :", len(secs))

    # --- Emplois (famille déduite du code) ---
    emps: dict[str, int] = {}
    for r in lire(dossier / "Emplois.xlsx"):
        code = str(r["code"])
        famille = FAMILLE_PAR_CODE.get(code)
        if famille is None:
            print("  ! emploi " + code + " : famille inconnue -> 3 (Employés-Ouvriers) par défaut")
            famille = 3
        e = Emploi(code=code, libelle=str(r["libelle"]), famille=famille)
        db.add(e); db.flush()
        emps[code] = e.id
    print("Emplois :", len(emps))

    # --- Catégories / Postes ---
    cats: dict[str, int] = {}
    for r in lire(dossier / "Categories.xlsx"):
        c = Categorie(code=str(r["code"]), libelle=str(r["code"]))
        db.add(c); db.flush()
        cats[c.code] = c.id
    print("Categories :", len(cats))

    posts: dict[str, int] = {}
    for r in lire(dossier / "Postes.xlsx"):
        p = Poste(code=str(r["code"]), libelle=str(r["libelle"]))
        db.add(p); db.flush()
        posts[p.code] = p.id
    print("Postes :", len(posts))

    # --- Critères (num legacy -> id) + détails (valeur = note cachée, OBLIGATOIRE) ---
    crits: dict[int, int] = {}
    for r in lire(dossier / "Criteres.xlsx"):
        c = Critere(code=str(r["num"]), libelle=str(r["libelle"]))
        db.add(c); db.flush()
        crits[int(r["num"])] = c.id
    print("Criteres :", len(crits))

    dets: dict[int, int] = {}
    for r in lire(dossier / "CritereDetails.xlsx"):
        if r["valeur"] is None or r["critere_num"] is None:
            print("  ! detail " + str(r["num"]) + " : critere_num/valeur MANQUANT -> "
                  "ré-extraire DetailCriteres avec (ID, CritereID, Libelle, Valeur)")
            continue
        critere_id = crits.get(int(r["critere_num"]))
        if not critere_id:
            print("  ! detail " + str(r["num"]) + " : critere " + str(r["critere_num"]) + " inconnu")
            continue
        d = CritereDetail(critere_id=critere_id,
                          libelle_descriptif=str(r["libelle"]),
                          valeur=float(r["valeur"]))
        db.add(d); db.flush()
        dets[int(r["num"])] = d.id
    print("CritereDetails :", len(dets))

    # --- Profils (num legacy -> id) ---
    profs: dict[int, int] = {}
    for r in lire(dossier / "Profils.xlsx"):
        p = Profil(code="P" + str(r["num"]), libelle=str(r["libelle"]))
        db.add(p); db.flush()
        profs[int(r["num"])] = p.id
    print("Profils :", len(profs))

    # --- ProfilCriteres (profil_num + critere_num + COEFF) ---
    nb_pc = 0
    for r in lire(dossier / "ProfilCriteres.xlsx"):
        pid = profs.get(int(r["profil_num"]))
        cid = crits.get(int(r["critere_num"]))
        if not pid or not cid:
            print("  ! profilcritere ignore : profil=" + str(r["profil_num"])
                  + " critere=" + str(r["critere_num"]))
            continue
        db.add(ProfilCritere(profil_id=pid, critere_id=cid,
                             coefficient=float(r["coefficient"]),
                             ordre=int(r["ordre"] or 0)))
        nb_pc += 1
    print("ProfilCriteres :", nb_pc)

    # --- EmploiProfils (emploi_code + profil_num) ---
    nb_ep = 0
    for r in lire(dossier / "EmploiProfils.xlsx"):
        eid = emps.get(str(r["emploi_code"]))
        pid = profs.get(int(r["profil_num"]))
        if not eid or not pid:
            print("  ! emploiprofil ignore : emploi=" + str(r["emploi_code"])
                  + " profil=" + str(r["profil_num"]))
            continue
        db.add(EmploiProfil(emploi_id=eid, profil_id=pid, ordre=int(r["ordre"] or 0)))
        nb_ep += 1
    print("EmploiProfils :", nb_ep)

    # --- Salaries (rattachements par code, N+1 par matricule) ---
    db.flush()
    matricules: dict[str, int] = {}
    rows_sal = lire(dossier / "Salaries.xlsx")
    # 1re passe : créer tous les salariés (sans N+1) pour avoir les matricules
    for r in rows_sal:
        mat = str(r["matricule"])
        if mat in matricules:
            print("  ! matricule dupliqué ignoré : " + mat)
            continue
        s = Salarie(
            matricule=mat, nom=str(r["nom"]), prenoms=r["prenoms"],
            site_id=sites.get(r["site_code"]) if r["site_code"] else None,
            departement_id=deps.get(r["dept_code"]) if r["dept_code"] else None,
            section_id=secs.get(r["section_code"]) if r["section_code"] else None,
            emploi_id=emps.get(r["emploi_code"]) if r["emploi_code"] else None,
            categorie_id=cats.get(r["categorie_code"]) if r["categorie_code"] else None,
            poste_id=posts.get(r["poste_code"]) if r["poste_code"] else None,
            date_embauche=r["date_embauche"],
            email=r["email"],
            hors_evaluation=en_bool(r["hors_evaluation"]),
            must_change_password=True,
            password_hash=hash_password(mat + "!prov"),
        )
        db.add(s); db.flush()
        matricules[mat] = s.id
    # 2e passe : lier les N+1 par matricule
    nb_n1 = 0
    for r in rows_sal:
        mat = str(r["matricule"])
        n1_mat = r["n1_matricule"]
        if not n1_mat:
            continue
        sid = matricules.get(mat)
        n1_id = matricules.get(str(n1_mat))
        if sid and n1_id:
            db.query(Salarie).filter(Salarie.id == sid).update({"n1_id": n1_id})
            nb_n1 += 1
        else:
            print("  ! N+1 non trouvé pour " + mat + " : " + str(n1_mat))
    print("Salaries :", len(matricules), "— liens N+1 :", nb_n1)

    db.commit(); db.close()
    print("✔ Référentiel importé. Mot de passe provisoire = matricule + !prov (ex. 3318!prov).")


if __name__ == "__main__":
    main()
