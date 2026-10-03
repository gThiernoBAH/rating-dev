#!/usr/bin/env python3
# 2026-10-02 — import Excel du référentiel (migration SQL Server -> openpyxl).
# Usage : python scripts/import_referentiel.py fichier.xlsx
# Feuilles attendues (une par onglet, en-têtes en ligne 1) :
#   Salaries, Sites, Departements, Sections, Emplois, Categories, Postes,
#   Criteres, CritereDetails, Profils, ProfilCriteres, EmploiProfils
import sys

sys.path.insert(0, ".")

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy.orm import Session   # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.models.rating import (  # noqa: E402
    Campagne, Categorie, Critere, CritereDetail, Departement, Emploi,
    EmploiProfil, Poste, Profil, ProfilCritere, Salarie, Section, Site,
)


def rows(ws):
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h else "" for h in next(it)]
    for r in it:
        if r and any(v is not None for v in r):
            yield dict(zip(headers, r))


def main() -> None:
    if len(sys.argv) != 2:
        print("Usage : import_referentiel.py fichier.xlsx")
        sys.exit(1)
    wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
    db: Session = SessionLocal()

    for ws in wb.worksheets:
        name = ws.title.strip()
        for row in rows(ws):
            if name == "Sites":
                db.add(Site(code=str(row["code"]), libelle=str(row["libelle"])))
            elif name == "Departements":
                db.add(Departement(code=str(row["code"]), libelle=str(row["libelle"]),
                                   site_id=int(row["site_id"])))
            elif name == "Sections":
                db.add(Section(code=str(row["code"]), libelle=str(row["libelle"]),
                               departement_id=int(row["departement_id"])))
            elif name == "Emplois":
                db.add(Emploi(code=str(row["code"]), libelle=str(row["libelle"]),
                              famille=int(row["famille"])))
            elif name == "Categories":
                db.add(Categorie(code=str(row["code"]), libelle=str(row["libelle"])))
            elif name == "Postes":
                db.add(Poste(code=str(row["code"]), libelle=str(row["libelle"])))
            elif name == "Criteres":
                db.add(Critere(code=str(row["code"]), libelle=str(row["libelle"])))
            elif name == "CritereDetails":
                db.add(CritereDetail(critere_id=int(row["critere_id"]),
                                     libelle_descriptif=str(row["libelle_descriptif"]),
                                     valeur=float(row["valeur"]),
                                     ordre=int(row.get("ordre") or 0)))
            elif name == "Profils":
                db.add(Profil(code=str(row["code"]), libelle=str(row["libelle"])))
            elif name == "ProfilCriteres":
                db.add(ProfilCritere(profil_id=int(row["profil_id"]),
                                     critere_id=int(row["critere_id"]),
                                     coefficient=float(row["coefficient"]),
                                     ordre=int(row.get("ordre") or 0)))
            elif name == "EmploiProfils":
                db.add(EmploiProfil(emploi_id=int(row["emploi_id"]),
                                    profil_id=int(row["profil_id"]),
                                    ordre=int(row.get("ordre") or 0)))
            elif name == "Salaries":
                # mot de passe provisoire imposé à la migration (§13.4)
                db.add(Salarie(
                    matricule=str(row["matricule"]), nom=str(row["nom"]),
                    prenoms=row.get("prenoms"),
                    site_id=row.get("site_id"), departement_id=row.get("departement_id"),
                    section_id=row.get("section_id"), emploi_id=row.get("emploi_id"),
                    categorie_id=row.get("categorie_id"), poste_id=row.get("poste_id"),
                    date_embauche=row.get("date_embauche"), email=row.get("email"),
                    n1_id=row.get("n1_id"),
                    hors_evaluation=bool(row.get("hors_evaluation") or False),
                    is_admin=bool(row.get("is_admin") or False),
                    must_change_password=True,
                    password_hash=hash_password(str(row["matricule"]) + "!prov"),
                ))
        print(f"  feuille {name} importée")

    db.commit()
    db.close()
    print("✔ Référentiel importé. Chaque salarié doit changer son mot de passe (provisoire = matricule!prov).")


if __name__ == "__main__":
    main()
